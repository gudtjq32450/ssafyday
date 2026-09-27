# -*- coding: utf-8 -*-
"""
SSAFY 싸피데이 인물퀴즈 MC 진행용 엑셀 정답표 생성 스크립트
index.html 내부의 300문제 데이터를 파싱하여 진행자(MC)가 실시간으로
정답/오답을 즉시 체크하고 진행할 수 있도록 최적화된 엑셀 문서를 생성합니다.
"""

import os
import sys
import json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding='utf-8')

BASE_DIR = r"C:\Users\wognl\OneDrive\바탕 화면\윤형섭\SSAFY관련\싸피데이용 게임\인물퀴즈"
HTML_PATH = os.path.join(BASE_DIR, "index.html")
OUTPUT_EXCEL = os.path.join(BASE_DIR, "인물퀴즈_MC진행용_정답표.xlsx")

def load_quizzes_from_html(html_path):
    with open(html_path, 'r', encoding='utf-8') as f:
        text = f.read()

    start_idx = text.find('createInitialQuizzes()')
    if start_idx == -1:
        raise ValueError("createInitialQuizzes() not found in index.html")
    
    array_start = text.find('[', start_idx)
    bracket_depth = 0
    array_end = -1
    in_string = False
    escape = False
    quote_char = None

    for i in range(array_start, len(text)):
        c = text[i]
        if escape:
            escape = False
            continue
        if c == '\\':
            escape = True
            continue
        if in_string:
            if c == quote_char:
                in_string = False
            continue
        else:
            if c in ('"', "'"):
                in_string = True
                quote_char = c
                continue
            if c == '[':
                bracket_depth += 1
            elif c == ']':
                bracket_depth -= 1
                if bracket_depth == 0:
                    array_end = i + 1
                    break

    quizzes = json.loads(text[array_start:array_end])
    return quizzes

def create_excel():
    print(f">>> index.html 로드 및 퀴즈 데이터 파싱 중: {HTML_PATH}")
    quizzes = load_quizzes_from_html(HTML_PATH)
    print(f">>> 총 {len(quizzes)}개 퀴즈 파싱 완료!")

    # 카테고리 정의 (정렬 및 스타일링용)
    CATEGORIES = [
        {
            "id": "anime",
            "name": "애니 & 캐릭터",
            "short_name": "애니&캐릭터",
            "main_color": "5B21B6",       # 진한 보라
            "sub_colors": ["7C3AED", "8B5CF6", "A78BFA"],
            "light_bg": ["F5F3FF", "EDE9FE"],
            "border_color": "C4B5FD",
            "tab_color": "8B5CF6",
            "levels": [
                {"lvl": 1, "pts": 10, "label": "1단계 (쉬움 / 10점)"},
                {"lvl": 2, "pts": 20, "label": "2단계 (보통 / 20점)"},
                {"lvl": 3, "pts": 30, "label": "3단계 (역전 / 30점)"}
            ]
        },
        {
            "id": "actor",
            "name": "배우 & 방송인",
            "short_name": "배우&방송인",
            "main_color": "1E40AF",       # 진한 블루
            "sub_colors": ["2563EB", "3B82F6", "60A5FA"],
            "light_bg": ["EFF6FF", "DBEAFE"],
            "border_color": "93C5FD",
            "tab_color": "3B82F6",
            "levels": [
                {"lvl": 1, "pts": 10, "label": "1단계 (쉬움 / 10점)"},
                {"lvl": 2, "pts": 20, "label": "2단계 (보통 / 20점)"},
                {"lvl": 3, "pts": 30, "label": "3단계 (역전 / 30점)"}
            ]
        },
        {
            "id": "music",
            "name": "가수 & 아이돌",
            "short_name": "가수&아이돌",
            "main_color": "065F46",       # 진한 에메랄드
            "sub_colors": ["059669", "10B981", "34D399"],
            "light_bg": ["ECFDF5", "D1FAE5"],
            "border_color": "6EE7B7",
            "tab_color": "10B981",
            "levels": [
                {"lvl": 1, "pts": 10, "label": "1단계 (쉬움 / 10점)"},
                {"lvl": 2, "pts": 20, "label": "2단계 (보통 / 20점)"},
                {"lvl": 3, "pts": 30, "label": "3단계 (역전 / 30점)"}
            ]
        },
        {
            "id": "youtube",
            "name": "유튜버 & 크리에이터",
            "short_name": "유튜버&크리에이터",
            "main_color": "9A3412",       # 진한 오렌지/레드
            "sub_colors": ["C2410C", "EA580C", "F97316"],
            "light_bg": ["FFF7ED", "FFEDD5"],
            "border_color": "FDBA74",
            "tab_color": "F97316",
            "levels": [
                {"lvl": 1, "pts": 10, "label": "1단계 (쉬움 / 10점)"},
                {"lvl": 2, "pts": 20, "label": "2단계 (보통 / 20점)"},
                {"lvl": 3, "pts": 30, "label": "3단계 (역전 / 30점)"}
            ]
        },
        {
            "id": "general",
            "name": "종합 & 일반인 [보너스]",
            "short_name": "종합&일반(보너스)",
            "main_color": "854D0E",       # 진한 골드/앰버
            "sub_colors": ["B45309", "D97706", "F59E0B"],
            "light_bg": ["FEFCE8", "FEF3C7"],
            "border_color": "FCD34D",
            "tab_color": "EAB308",
            "levels": [
                {"lvl": 1, "pts": 15, "label": "1단계 (쉬움 / 15점)"},
                {"lvl": 2, "pts": 25, "label": "2단계 (보통 / 25점)"},
                {"lvl": 3, "pts": 40, "label": "3단계 (역전 / 40점)"}
            ]
        }
    ]

    # roundKey별로 묶기
    round_dict = {}
    for q in quizzes:
        rk = q['roundKey']
        if rk not in round_dict:
            round_dict[rk] = []
        round_dict[rk].append(q)
    
    # qNum 순서대로 정렬 확인
    for rk in round_dict:
        round_dict[rk].sort(key=lambda x: x['qNum'])

    # 스타일 공통 정의
    font_family = "맑은 고딕"
    
    thin_side = Side(border_style="thin", color="CBD5E1")
    medium_side = Side(border_style="medium", color="64748B")
    double_side = Side(border_style="double", color="334155")
    
    cell_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
    thick_bottom_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=medium_side)

    wb = openpyxl.Workbook()
    # 기본 첫 시트 제거 또는 재사용
    ws_board = wb.active
    ws_board.title = "🎯 한눈에보기(15개라운드)"
    ws_board.sheet_properties.tabColor = "DC2626" # 눈에 띄는 빨간 탭

    # =========================================================================
    # [시트 1] 🎯 [한눈에보기] 15개라운드 정답판 (MC 실시간 현장 진행 최적화)
    # =========================================================================
    print(">>> [시트 1] 15개 라운드 한눈에 보는 보드판 작성 중...")
    
    # 1. 헤더 타이틀 배너
    ws_board.merge_cells("A1:P1")
    title_cell = ws_board["A1"]
    title_cell.value = "🎯 SSAFY 싸피데이 인물 맞추기 게임 — MC 실시간 진행용 정답표 (15개 라운드 총 300문제)"
    title_cell.font = Font(name=font_family, size=15, bold=True, color="FFFFFF")
    title_cell.fill = PatternFill(fill_type="solid", start_color="0F172A", end_color="0F172A")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws_board.row_dimensions[1].height = 42

    # 2. 설명 배너
    ws_board.merge_cells("A2:P2")
    info_cell = ws_board["A2"]
    info_cell.value = "💡 [진행 꿀팁] 스태프가 Space를 눌러 문제를 열면, MC는 해당 열(카테고리/단계)의 Q.1~Q.20을 보며 정답 여부를 즉시 판정하세요! (Space: 정답 인정 / X: 땡&토스 / →: 패스)"
    info_cell.font = Font(name=font_family, size=10, bold=True, color="1E293B")
    info_cell.fill = PatternFill(fill_type="solid", start_color="F1F5F9", end_color="F1F5F9")
    info_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws_board.row_dimensions[2].height = 26

    # 3. 빈 줄
    ws_board.row_dimensions[3].height = 10

    # 4. 카테고리 대분류 헤더 (행 4)
    ws_board.cell(row=4, column=1, value="구분").fill = PatternFill(fill_type="solid", start_color="334155", end_color="334155")
    ws_board.cell(row=4, column=1).font = Font(name=font_family, size=11, bold=True, color="FFFFFF")
    ws_board.cell(row=4, column=1).alignment = Alignment(horizontal="center", vertical="center")
    ws_board.cell(row=4, column=1).border = cell_border

    col_idx = 2
    for cat in CATEGORIES:
        start_col = col_idx
        end_col = col_idx + 2
        ws_board.merge_cells(start_row=4, start_column=start_col, end_row=4, end_column=end_col)
        h_cell = ws_board.cell(row=4, column=start_col)
        h_cell.value = f"{cat['name']}"
        h_cell.font = Font(name=font_family, size=12, bold=True, color="FFFFFF")
        h_cell.fill = PatternFill(fill_type="solid", start_color=cat["main_color"], end_color=cat["main_color"])
        h_cell.alignment = Alignment(horizontal="center", vertical="center")
        
        for c in range(start_col, end_col + 1):
            ws_board.cell(row=4, column=c).border = cell_border
        
        col_idx += 3
    ws_board.row_dimensions[4].height = 30

    # 5. 난이도/배점 소분류 헤더 (행 5)
    ws_board.cell(row=5, column=1, value="문제 No.").fill = PatternFill(fill_type="solid", start_color="475569", end_color="475569")
    ws_board.cell(row=5, column=1).font = Font(name=font_family, size=10, bold=True, color="FFFFFF")
    ws_board.cell(row=5, column=1).alignment = Alignment(horizontal="center", vertical="center")
    ws_board.cell(row=5, column=1).border = thick_bottom_border

    col_idx = 2
    round_columns = [] # (rk, col_idx, cat, lvl_info)
    for cat in CATEGORIES:
        for i, lvl_info in enumerate(cat["levels"]):
            rk = f"{cat['id']}_l{lvl_info['lvl']}"
            round_columns.append((rk, col_idx, cat, lvl_info))
            sub_cell = ws_board.cell(row=5, column=col_idx)
            sub_cell.value = f"{lvl_info['lvl']}단계 ({lvl_info['pts']}점)"
            sub_cell.font = Font(name=font_family, size=10, bold=True, color="FFFFFF")
            sub_cell.fill = PatternFill(fill_type="solid", start_color=cat["sub_colors"][i], end_color=cat["sub_colors"][i])
            sub_cell.alignment = Alignment(horizontal="center", vertical="center")
            sub_cell.border = thick_bottom_border
            col_idx += 1
    ws_board.row_dimensions[5].height = 26

    # 6. 문제 1번 ~ 20번 데이터 행 채우기 (행 6 ~ 25)
    for q_idx in range(1, 21):
        row_num = 5 + q_idx
        ws_board.row_dimensions[row_num].height = 28
        
        # A열: 문제 번호 Q.01 ~ Q.20
        q_label_cell = ws_board.cell(row=row_num, column=1, value=f"Q.{q_idx:02d}")
        q_label_cell.font = Font(name=font_family, size=11, bold=True, color="0F172A")
        q_label_cell.alignment = Alignment(horizontal="center", vertical="center")
        q_label_cell.fill = PatternFill(fill_type="solid", start_color="F1F5F9" if q_idx % 2 == 1 else "E2E8F0", 
                                        end_color="F1F5F9" if q_idx % 2 == 1 else "E2E8F0")
        q_label_cell.border = cell_border

        # 15개 열 정답 데이터 채우기
        for rk, c_idx, cat, lvl_info in round_columns:
            q_list = round_dict.get(rk, [])
            ans_text = ""
            if q_idx <= len(q_list):
                ans_text = q_list[q_idx - 1]["answer"]

            ans_cell = ws_board.cell(row=row_num, column=c_idx, value=ans_text)
            ans_cell.font = Font(name=font_family, size=11, bold=True, color="0F172A")
            ans_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            
            # 교차 배경색 (가독성 극대화)
            bg_color = cat["light_bg"][0] if q_idx % 2 == 1 else "FFFFFF"
            ans_cell.fill = PatternFill(fill_type="solid", start_color=bg_color, end_color=bg_color)
            ans_cell.border = cell_border

    # 열 너비 설정
    ws_board.column_dimensions["A"].width = 10
    for c in range(2, 17):
        col_letter = get_column_letter(c)
        ws_board.column_dimensions[col_letter].width = 17

    # 틀 고정: Q번호와 헤더 고정 (B6에서 고정)
    ws_board.freeze_panes = "B6"

    # 눈금선 표시 및 A4 가로 인쇄 맞춤 설정
    ws_board.views.sheetView[0].showGridLines = True
    ws_board.page_setup.orientation = ws_board.ORIENTATION_LANDSCAPE
    ws_board.page_setup.paperSize = ws_board.PAPERSIZE_A4
    ws_board.page_setup.fitToPage = True
    ws_board.page_setup.fitToWidth = 1
    ws_board.page_setup.fitToHeight = 1
    ws_board.sheet_properties.pageSetUpPr.fitToPage = True

    # =========================================================================
    # [시트 2] 📋 전체_300문제_상세정답목록 (필터 및 검색용)
    # =========================================================================
    print(">>> [시트 2] 300문제 전체 상세 정답 목록 작성 중...")
    ws_all = wb.create_sheet(title="📋 전체상세(300문제)")
    ws_all.sheet_properties.tabColor = "2563EB"

    headers_all = [
        ("No.", 7, "center"),
        ("카테고리", 18, "center"),
        ("난이도", 10, "center"),
        ("기본배점", 10, "center"),
        ("세트내 번호", 12, "center"),
        ("정답 (인물 / 캐릭터명)", 26, "left"),
        ("힌트 및 작품/활동 배경", 50, "left"),
        ("이미지 파일 경로", 35, "left")
    ]

    # 헤더 작성
    ws_all.row_dimensions[1].height = 32
    for c_idx, (h_title, width, align) in enumerate(headers_all, start=1):
        cell = ws_all.cell(row=1, column=c_idx, value=h_title)
        cell.font = Font(name=font_family, size=11, bold=True, color="FFFFFF")
        cell.fill = PatternFill(fill_type="solid", start_color="1E293B", end_color="1E293B")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thick_bottom_border
        col_letter = get_column_letter(c_idx)
        ws_all.column_dimensions[col_letter].width = width

    # 카테고리 맵
    cat_map = {c["id"]: c for c in CATEGORIES}

    # 데이터 작성
    for idx, q in enumerate(quizzes, start=1):
        row_num = idx + 1
        ws_all.row_dimensions[row_num].height = 24
        
        cat_info = cat_map.get(q["catId"], {})
        cat_name = q.get("catName", "")
        lvl_str = f"{q.get('level')}단계"
        pts_str = f"{q.get('points')}점"
        q_num_str = f"Q.{q.get('qNum'):02d}"
        ans = q.get("answer", "")
        hint = q.get("hint", "")
        img_url = q.get("imageUrl", "")

        row_values = [idx, cat_name, lvl_str, pts_str, q_num_str, ans, hint, img_url]
        
        bg_color = "F8FAFC" if idx % 2 == 1 else "FFFFFF"

        for c_idx, val in enumerate(row_values, start=1):
            cell = ws_all.cell(row=row_num, column=c_idx, value=val)
            align = headers_all[c_idx - 1][2]
            cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=(c_idx in [6, 7]))
            cell.border = cell_border
            cell.fill = PatternFill(fill_type="solid", start_color=bg_color, end_color=bg_color)
            
            if c_idx == 6:  # 정답 열 강조
                cell.font = Font(name=font_family, size=11, bold=True, color="1E3A8A")
            elif c_idx == 1 or c_idx == 5:
                cell.font = Font(name=font_family, size=10, bold=True, color="475569")
            else:
                cell.font = Font(name=font_family, size=10, color="0F172A")

    # 필터 적용 및 틀 고정
    ws_all.auto_filter.ref = f"A1:H{len(quizzes) + 1}"
    ws_all.freeze_panes = "A2"
    ws_all.views.sheetView[0].showGridLines = True

    # =========================================================================
    # [시트 3~7] 5대 개별 카테고리 상세 시트 (1/2/3단계 나란히 배치)
    # =========================================================================
    for cat_idx, cat in enumerate(CATEGORIES, start=1):
        cat_title = f"{cat_idx}. {cat['short_name']}"
        print(f">>> [시트 {cat_idx + 2}] 카테고리 '{cat['name']}' 상세 시트 작성 중...")
        ws_cat = wb.create_sheet(title=cat_title)
        ws_cat.sheet_properties.tabColor = cat["tab_color"]

        # 1. 상단 타이틀
        ws_cat.merge_cells("A1:K1")
        t_cell = ws_cat["A1"]
        t_cell.value = f"📌 {cat['name']} — 난이도별 정답 및 힌트 상세 일람표"
        t_cell.font = Font(name=font_family, size=14, bold=True, color="FFFFFF")
        t_cell.fill = PatternFill(fill_type="solid", start_color=cat["main_color"], end_color=cat["main_color"])
        t_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws_cat.row_dimensions[1].height = 38

        # 2. 난이도별 3개 블록 구성:
        # 블록 1 (1단계): A~C열 (A: 번호, B: 정답, C: 힌트/상세)
        # 구분선: D열 (빈 공백)
        # 블록 2 (2단계): E~G열 (E: 번호, F: 정답, G: 힌트/상세)
        # 구분선: H열 (빈 공백)
        # 블록 3 (3단계): I~K열 (I: 번호, J: 정답, K: 힌트/상세)

        block_starts = [1, 5, 9] # A, E, I 열
        for i, lvl_info in enumerate(cat["levels"]):
            b_start = block_starts[i]
            b_end = b_start + 2
            
            # 단계 헤더 (행 2)
            ws_cat.merge_cells(start_row=2, start_column=b_start, end_row=2, end_column=b_end)
            h_lvl = ws_cat.cell(row=2, column=b_start)
            h_lvl.value = f"⭐ {lvl_info['label']}"
            h_lvl.font = Font(name=font_family, size=11, bold=True, color="FFFFFF")
            h_lvl.fill = PatternFill(fill_type="solid", start_color=cat["sub_colors"][i], end_color=cat["sub_colors"][i])
            h_lvl.alignment = Alignment(horizontal="center", vertical="center")
            
            for c in range(b_start, b_end + 1):
                ws_cat.cell(row=2, column=c).border = cell_border

            # 서브 헤더 (행 3)
            ws_cat.cell(row=3, column=b_start, value="No.").alignment = Alignment(horizontal="center", vertical="center")
            ws_cat.cell(row=3, column=b_start + 1, value="정답 (인물명)").alignment = Alignment(horizontal="center", vertical="center")
            ws_cat.cell(row=3, column=b_start + 2, value="힌트 / 상세 설명").alignment = Alignment(horizontal="center", vertical="center")
            
            for c in range(b_start, b_end + 1):
                c_cell = ws_cat.cell(row=3, column=c)
                c_cell.font = Font(name=font_family, size=10, bold=True, color="FFFFFF")
                c_cell.fill = PatternFill(fill_type="solid", start_color="475569", end_color="475569")
                c_cell.border = thick_bottom_border

            # 20문제 데이터 작성 (행 4 ~ 23)
            rk = f"{cat['id']}_l{lvl_info['lvl']}"
            q_list = round_dict.get(rk, [])
            
            for q_idx in range(1, 21):
                r_num = 3 + q_idx
                ws_cat.row_dimensions[r_num].height = 26
                
                q_data = q_list[q_idx - 1] if q_idx <= len(q_list) else {}
                ans_str = q_data.get("answer", "")
                hint_str = q_data.get("hint", "")

                bg_c = cat["light_bg"][0] if q_idx % 2 == 1 else "FFFFFF"

                # 번호
                c_num = ws_cat.cell(row=r_num, column=b_start, value=f"Q.{q_idx:02d}")
                c_num.font = Font(name=font_family, size=10, bold=True, color="334155")
                c_num.alignment = Alignment(horizontal="center", vertical="center")
                c_num.fill = PatternFill(fill_type="solid", start_color=bg_c, end_color=bg_c)
                c_num.border = cell_border

                # 정답
                c_ans = ws_cat.cell(row=r_num, column=b_start + 1, value=ans_str)
                c_ans.font = Font(name=font_family, size=11, bold=True, color="0F172A")
                c_ans.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                c_ans.fill = PatternFill(fill_type="solid", start_color=bg_c, end_color=bg_c)
                c_ans.border = cell_border

                # 힌트/설명
                c_hint = ws_cat.cell(row=r_num, column=b_start + 2, value=hint_str)
                c_hint.font = Font(name=font_family, size=9, color="475569")
                c_hint.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
                c_hint.fill = PatternFill(fill_type="solid", start_color=bg_c, end_color=bg_c)
                c_hint.border = cell_border

        ws_cat.row_dimensions[2].height = 26
        ws_cat.row_dimensions[3].height = 24

        # 열 너비 설정
        ws_cat.column_dimensions["A"].width = 8
        ws_cat.column_dimensions["B"].width = 18
        ws_cat.column_dimensions["C"].width = 28
        ws_cat.column_dimensions["D"].width = 3   # 구분 공백열

        ws_cat.column_dimensions["E"].width = 8
        ws_cat.column_dimensions["F"].width = 18
        ws_cat.column_dimensions["G"].width = 28
        ws_cat.column_dimensions["H"].width = 3   # 구분 공백열

        ws_cat.column_dimensions["I"].width = 8
        ws_cat.column_dimensions["J"].width = 18
        ws_cat.column_dimensions["K"].width = 28

        # 틀 고정 및 눈금선
        ws_cat.freeze_panes = "A4"
        ws_cat.views.sheetView[0].showGridLines = True

    # =========================================================================
    # [시트 8] 💡 MC_진행가이드&치트시트
    # =========================================================================
    print(">>> [시트 8] MC 진행 가이드 & 치트시트 시트 작성 중...")
    ws_guide = wb.create_sheet(title="💡 MC_진행가이드&단축키")
    ws_guide.sheet_properties.tabColor = "475569"

    guide_content = [
        ("🎯 SSAFY 인물 맞추기 게임 MC 실시간 진행 & 판정 매뉴얼", "title"),
        ("", "blank"),
        ("1. 조작 담당 스태프와의 찰떡 호흡 요령 (진행자 핵심 가이드)", "h1"),
        ("• 상황 요약: 프로젝터 화면을 조작하는 컴퓨터는 스태프가 누르고, MC는 마이크를 쥐고 현장을 리드합니다.", "bullet"),
        ("• Step 1. 문제 출제 준비: 화면에 [O조 OOO 타자] 차례가 뜨면 참가자가 준비되었는지 확인 후 MC가 '문제 나갑니다!' 외침", "bullet"),
        ("• Step 2. 사진 공개: 스태프가 [Space] 키를 누르면 사진이 온전히 열립니다. MC는 마이크로 즉시 구두 카운트다운! ('하나, 둘, 셋!')", "bullet"),
        ("• Step 3. 정답 판정 (맞혔을 때): 도전자 답변이 맞으면 MC가 씩씩하게 '정답!' 외침 ➔ 스태프가 [Space] 키를 누름", "bullet"),
        ("  - ➔ 화면에 정답 배너가 뜨며 점수가 자동 누적되고, 3초 카운트다운 후 자동으로 다음 문제로 이동합니다.", "subbullet"),
        ("• Step 4. 실패 판정 (틀렸을 때): 도전자 답변이 틀리거나 3초 초과 시 MC가 '땡!' 외침 ➔ 스태프가 [X] 키를 누름", "bullet"),
        ("  - ➔ 사진이 닫히지 않고 그대로 노출된 상태에서 즉시 다음 순번 조 도전자에게 기회가 토스됩니다! (+2점 가점 누적!)", "subbullet"),
        ("• Step 5. 전원 실패 (패스): 4개 조가 한 바퀴 돌며 모두 실패한 경우 ➔ 스태프가 [→] 또는 [N]을 누르면 점수 획득 없이 정답 공개 후 다음 문제 이동", "bullet"),
        ("", "blank"),
        ("2. 컴퓨터 조작 스태프 키보드 단축키 치트시트", "h1"),
        ("• [Space] (스피디 원키) : (1) 가림막 열고 사진 공개 ➔ (2) 정답 인정 & 점수 반영 및 다음 문제 카운트다운", "bullet"),
        ("• [X] 또는 [ㅌ] : 땡! (실패 & 즉시 기회 토스, 누적 가점 +2점)", "bullet"),
        ("• [→] 또는 [N] : 패스 (4개 조 모두 실패 시 점수 없이 정답 공개 후 다음 문제로 스킵)", "bullet"),
        ("• [Esc] : 진행 중 메인 보드판으로 즉시 복귀", "bullet"),
        ("• [B] 또는 [ㅠ] : 긴장감 넘치는 배경음악(BGM) On/Off 토글", "bullet"),
        ("• 상단 점수 클릭 : 진행 도중 점수 오차가 있을 경우 숫자를 직접 클릭하여 자유롭게 타이핑 수정 가능!", "bullet"),
        ("", "blank"),
        ("3. MC 유연한 정답 인정 기준 가이드라인 (현장 실전 팁)", "h1"),
        ("• 본명 vs 예명: '감스트' vs '김인직', '덱스' vs '김진영', '침착맨' vs '이말년' 등 대중적으로 널리 알려진 이름은 모두 정답 인정!", "bullet"),
        ("• 애니메이션 캐릭터: '디바' or '송하나', '골롬보 반장' or '메구레 쥬조', '괴도 키드' or '쿠로바 카이토' 등 캐릭터명과 극중 본명 모두 인정!", "bullet"),
        ("• 성 생략 여부: 외국인 이름의 경우 통상 불리는 이름('메시', '호날두', '사쿠라') 인정 가능 / 한국인은 풀네임 3글자 원칙 권장!", "bullet"),
        ("• 3단계 함정 문제 주의: '궁극의 푸른눈의 백룡'은 그냥 '푸른눈의 백룡'이라고 하면 땡! '궁극의'까지 외쳐야 정답 인정!", "bullet"),
        ("• 싸피 강사/프로님 문제: 종합 3단계 등에 포함된 반 강사님/프로님 문제는 성함을 정확히 맞혀야 인정!", "bullet"),
        ("", "blank"),
        ("4. 세트 구성 및 다음 게임 선택권 룰", "h1"),
        ("• 총 6개 게임 세트 진행 (각 세트 20문제 연속 릴레이)", "bullet"),
        ("• 1~5게임은 1~5번 타자 출전 / 6게임(파이널)은 6인 조는 6번 타자, 5인 조는 1번 타자 재출전!", "bullet"),
        ("• 매 세트 마지막 20번째 문제를 맞힌 사람이 [다음 게임의 카테고리와 난이도를 선택하는 황금 선택권]을 획득합니다!", "bullet")
    ]

    ws_guide.column_dimensions["A"].width = 5
    ws_guide.column_dimensions["B"].width = 110

    for r_idx, (text_val, style_type) in enumerate(guide_content, start=1):
        ws_guide.row_dimensions[r_idx].height = 24
        cell = ws_guide.cell(row=r_idx, column=2, value=text_val)
        
        if style_type == "title":
            ws_guide.row_dimensions[r_idx].height = 36
            cell.font = Font(name=font_family, size=14, bold=True, color="FFFFFF")
            cell.fill = PatternFill(fill_type="solid", start_color="0F172A", end_color="0F172A")
            cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        elif style_type == "h1":
            ws_guide.row_dimensions[r_idx].height = 28
            cell.font = Font(name=font_family, size=11, bold=True, color="FFFFFF")
            cell.fill = PatternFill(fill_type="solid", start_color="1E3A8A", end_color="1E3A8A")
            cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        elif style_type == "bullet":
            cell.font = Font(name=font_family, size=10, bold=True, color="1E293B")
            cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        elif style_type == "subbullet":
            cell.font = Font(name=font_family, size=9, color="475569")
            cell.alignment = Alignment(horizontal="left", vertical="center", indent=2)

    # 엑셀 파일 저장 (인물퀴즈_MC진행용_정답표.xlsx 및 v2 파일 동시 최신화)
    try:
        wb.save(OUTPUT_EXCEL)
        print(f"\n🎉 [기본] 엑셀 정답표 저장 완료: {OUTPUT_EXCEL}")
    except Exception as e:
        print(f"⚠️ 기본 파일 저장 실패 ({e})")

    try:
        v2_path = os.path.join(BASE_DIR, "인물퀴즈_MC진행용_정답표_v2.xlsx")
        wb.save(v2_path)
        print(f"🎉 [v2] 엑셀 정답표 저장 완료: {v2_path}")
    except Exception as e:
        print(f"⚠️ v2 파일 저장 실패 ({e})")

if __name__ == "__main__":
    create_excel()
