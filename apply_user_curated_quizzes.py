# -*- coding: utf-8 -*-
"""
사용자 직접 검수 및 폴더 정리 완료된 301문제 index.html 전면 반영 &
카테고리 보드 15개 버튼 대표 인물/캐릭터 블러 배경 연동 스크립트
"""

import os
import io
import re
import sys
import json
import shutil
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')

BASE_DIR = r"C:\Users\wognl\OneDrive\바탕 화면\윤형섭\SSAFY관련\싸피데이용 게임\인물퀴즈"
IMAGES_DIR = os.path.join(BASE_DIR, "images")
EXCEL_FILE = os.path.join(BASE_DIR, "한국_2030_분야별_인지도_TOP100_총500인.xlsx")
HTML_FILE = os.path.join(BASE_DIR, "index.html")

# 엑셀 시트 매핑
SHEET_MAP = {
    "anime": "애니메이션 & 캐릭터",
    "actor": "배우 & 방송인",
    "music": "가수 & 아이돌",
    "youtube": "유튜버 & 크리에이터",
    "general": "스포츠·기업인·프로게이머"
}

# 15개 폴더 정의
FOLDERS = [
    ("애니_1단계", "anime", "애니 & 캐릭터", 1, 10),
    ("애니_2단계", "anime", "애니 & 캐릭터", 2, 20),
    ("애니_3단계", "anime", "애니 & 캐릭터", 3, 30),
    ("배우_1단계", "actor", "배우 & 방송인", 1, 10),
    ("배우_2단계", "actor", "배우 & 방송인", 2, 20),
    ("배우_3단계", "actor", "배우 & 방송인", 3, 30),
    ("가수_1단계", "music", "가수 & 아이돌", 1, 10),
    ("가수_2단계", "music", "가수 & 아이돌", 2, 20),
    ("가수_3단계", "music", "가수 & 아이돌", 3, 30),
    ("유튜버_1단계", "youtube", "유튜버 & 크리에이터", 1, 10),
    ("유튜버_2단계", "youtube", "유튜버 & 크리에이터", 2, 20),
    ("유튜버_3단계", "youtube", "유튜버 & 크리에이터", 3, 30),
    ("종합_1단계", "general", "종합 & 일반인", 1, 15),
    ("종합_2단계", "general", "종합 & 일반인", 2, 25),
    ("종합_3단계", "general", "종합 & 일반인", 3, 40)
]

# 1. 엑셀에서 카테고리별 힌트 딕셔너리 구축
print(">>> 엑셀 파일에서 카테고리별 힌트 데이터 로드...")
sheet_hints = {k: {} for k in SHEET_MAP.keys()}
try:
    wb = openpyxl.load_workbook(EXCEL_FILE, data_only=True)
    for cat_id, sname in SHEET_MAP.items():
        if sname in wb.sheetnames:
            sheet = wb[sname]
            for row in sheet.iter_rows(values_only=True):
                if row and len(row) > 2 and row[2]:
                    name = str(row[2]).strip()
                    detail = str(row[3]).strip() if len(row) > 3 and row[3] else ""
                    hint = str(row[5]).strip() if len(row) > 5 and row[5] else ""
                    h = hint or detail
                    if h:
                        sheet_hints[cat_id][name] = h
                        clean = re.sub(r'\(.*?\)', '', name).strip()
                        if clean: sheet_hints[cat_id][clean] = h
    # 엄태구 힌트 추가
    sheet_hints["actor"]["엄태구"] = "허스키한 동굴 보이스와 카리스마, 영화 <낙원의 밤>, 드라마 <놀아주는 여자>의 대세 배우"
except Exception as e:
    print(f"엑셀 로드 경고: {e}")

# 2. 15개 폴더의 실제 파일 스캔 및 퀴즈 객체 생성
print("\n>>> 사용자 검수 폴더에서 퀴즈 목록 생성...")
quizzes = []

for folder, cat_id, cat_title, lvl, pts in FOLDERS:
    folder_path = os.path.join(IMAGES_DIR, folder)
    if not os.path.exists(folder_path):
        print(f"경고: {folder_path} 폴더가 없습니다.")
        continue

    files = sorted([f for f in os.listdir(folder_path) if f.lower().endswith(('.jpg', '.jpeg', '.png', '.webp'))])
    print(f"  [{folder}] {len(files)}개 파일 로드")

    for q_idx, fname in enumerate(files, start=1):
        ans = os.path.splitext(fname)[0].strip()
        qid = f"{cat_id}_l{lvl}_q{q_idx}"
        rel_img_url = f"images/{folder}/{fname}"

        # 힌트 매칭
        hint = sheet_hints[cat_id].get(ans, "")
        if not hint:
            clean_ans = re.sub(r'\(.*?\)', '', ans).strip()
            hint = sheet_hints[cat_id].get(clean_ans, "")
        if not hint and "(" in ans and ")" in ans:
            # 괄호 안의 내용을 힌트로 활용
            m = re.findall(r'\((.*?)\)', ans)
            if m: hint = ", ".join(m)
        if not hint:
            hint = f"{cat_title} {lvl}단계"

        quizzes.append({
            "id": qid,
            "catId": cat_id,
            "catName": cat_title,
            "level": lvl,
            "qNum": q_idx,
            "roundKey": f"{cat_id}_l{lvl}",
            "points": pts,
            "answer": ans,
            "hint": hint,
            "imageUrl": rel_img_url
        })

        # fallback images/{id}.jpg 미러링
        alt_dest = os.path.join(IMAGES_DIR, f"{qid}.jpg")
        src_path = os.path.join(folder_path, fname)
        try:
            shutil.copyfile(src_path, alt_dest)
        except Exception:
            pass

print(f"\n총 생성된 문제 수: {len(quizzes)}개")

# 3. index.html 업데이트
print("\n>>> index.html 소스코드 최신화...")
with open(HTML_FILE, "r", encoding="utf-8") as f:
    html = f.read()

# 3-1. CSS 업데이트: .level-round-card 에 블러 배경 레이어 스타일 주입
css_target = """        /* 카테고리 단계별 20문제 라운드 카드 */
        .level-round-card {
            flex: 1;
            background: rgba(255, 255, 255, 0.07);
            border-radius: 12px;
            display: flex;
            flex-direction: column;
            align-items: center;
            justify-content: center;
            gap: 4px;
            cursor: pointer;
            transition: all 0.25s ease;
            position: relative;
            overflow: hidden;
            padding: 8px 6px;
        }"""

css_replacement = """        /* 카테고리 단계별 20문제 라운드 카드 (인물/캐릭터 블러 배경 적용) */
        .level-round-card {
            flex: 1;
            background: rgba(15, 20, 35, 0.7);
            border-radius: 12px;
            display: flex;
            align-items: center;
            justify-content: center;
            cursor: pointer;
            transition: all 0.25s cubic-bezier(0.4, 0, 0.2, 1);
            position: relative;
            overflow: hidden;
            padding: 8px 6px;
        }
        .level-round-card .card-bg-blur {
            position: absolute;
            inset: -4px;
            background-image: var(--btn-bg);
            background-size: cover;
            background-position: center 20%;
            filter: blur(2.5px) brightness(0.38) saturate(1.2);
            transform: scale(1.08);
            transition: all 0.3s ease;
            z-index: 1;
        }
        .level-round-card:hover:not(.solved) .card-bg-blur {
            filter: blur(1.5px) brightness(0.52) saturate(1.35);
            transform: scale(1.14);
        }
        .level-round-card .card-text-content {
            position: relative;
            z-index: 2;
            display: flex;
            flex-direction: column;
            align-items: center;
            justify-content: center;
            gap: 4px;
            width: 100%;
            pointer-events: none;
            text-shadow: 0 2px 8px rgba(0, 0, 0, 0.95), 0 0 4px #000;
        }"""

if css_target in html:
    html = html.replace(css_target, css_replacement)
    print("  CSS .level-round-card 스타일 교체 완료")
else:
    # 혹시 이미 수정되었거나 형식이 다를 경우 정규식 대치
    pattern = r"/\* 카테고리 단계별 20문제 라운드 카드[^*]*\*/\s*\.level-round-card\s*\{[^}]*\}"
    html = re.sub(pattern, css_replacement, html)
    print("  CSS .level-round-card 정규식 교체 완료")

# 3-2. solved 상태 스타일 보강
solved_css_target = """.level-round-card.solved {
            background: rgba(255, 42, 109, 0.12) !important;
            border-color: rgba(255, 42, 109, 0.6) !important;
            cursor: not-allowed !important;
            opacity: 0.65;
            transform: none !important;
            box-shadow: none !important;
        }"""

solved_css_replacement = """.level-round-card.solved {
            border-color: rgba(255, 42, 109, 0.5) !important;
            cursor: not-allowed !important;
            transform: none !important;
            box-shadow: none !important;
        }
        .level-round-card.solved .card-bg-blur {
            filter: blur(5px) grayscale(0.9) brightness(0.2) !important;
        }
        .level-round-card.solved .card-text-content {
            opacity: 0.4;
        }"""

if solved_css_target in html:
    html = html.replace(solved_css_target, solved_css_replacement)
    print("  CSS .level-round-card.solved 스타일 교체 완료")

# 3-3. renderBoard() 내부 카드 렌더링 로직 수정 (블러 배경 div 및 CSS 변수 바인딩)
old_render_block = """                    const card = document.createElement('div');
                    card.className = `level-round-card lvl-${lvl} ${isRoundDone ? 'solved' : ''}`;

                    card.innerHTML = `
                        <div class="lvl-title">⭐ ${lvl}단계</div>
                        <div class="lvl-sub">20문제 릴레이 (${solvedCount}/20)</div>
                        <div class="lvl-points">${ptsText}</div>
                    `;"""

new_render_block = """                    const card = document.createElement('div');
                    card.className = `level-round-card lvl-${lvl} ${isRoundDone ? 'solved' : ''}`;
                    const btnBgUrl = `images/board_buttons/${cat.id}_l${lvl}.jpg`;
                    card.style.setProperty('--btn-bg', `url('${btnBgUrl}')`);

                    card.innerHTML = `
                        <div class="card-bg-blur"></div>
                        <div class="card-text-content">
                            <div class="lvl-title">⭐ ${lvl}단계</div>
                            <div class="lvl-sub">${roundQuizzes.length}문제 릴레이 (${solvedCount}/${roundQuizzes.length})</div>
                            <div class="lvl-points">${ptsText}</div>
                        </div>
                    `;"""

if old_render_block in html:
    html = html.replace(old_render_block, new_render_block)
    print("  renderBoard() 내부 카드 마크업 교체 완료")
else:
    # 정규식 치환
    pattern = r"const card = document\.createElement\('div'\);\s*card\.className = `level-round-card lvl-\$\{lvl\} \$\{isRoundDone \? 'solved' : ''\}`;[\s\S]*?card\.innerHTML = `[\s\S]*?`;"
    html = re.sub(pattern, new_render_block, html)
    print("  renderBoard() 정규식 교체 완료")

# 3-4. createInitialQuizzes() 함수 교체
js_quizzes = json.dumps(quizzes, ensure_ascii=False, indent=8)
new_create_func = f"""function createInitialQuizzes() {{
            return {js_quizzes};
        }}"""

pattern = r"function createInitialQuizzes\(\)\s*\{[\s\S]*?return\s*\[[\s\S]*?\];\s*\}"
html = re.sub(pattern, new_create_func, html)
print("  createInitialQuizzes() 301문제 데이터 주입 완료")

# 3-5. LocalStorage 버전 v7 -> v8로 승격하여 브라우저 강제 갱신
for v in ['v2', 'v3', 'v4', 'v5', 'v6', 'v7']:
    html = html.replace(f'ssafy_quiz_list_{v}', 'ssafy_quiz_list_v8')
    html = html.replace(f'ssafy_solved_status_{v}', 'ssafy_solved_status_v8')
    html = html.replace(f'ssafy_round_status_{v}', 'ssafy_round_status_v8')

print("  LocalStorage 스토리지 버전 v8로 승격 완료")

# 파일 저장
with open(HTML_FILE, "w", encoding="utf-8") as f:
    f.write(html)

print("\n🎉 모든 작업이 성공적으로 완료되었습니다!")
